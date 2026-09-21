from datetime import datetime
from io import BytesIO
import re

from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    UploadFile,
    File,
)

from pydantic import BaseModel

from sqlalchemy.orm import Session

from openpyxl import load_workbook

from .database import get_db
from .models import (
    User,
    Course,
    Timetable,
    Enrollment,
)

from .dependencies import current_user, audit


# ============================================================
# ROUTERS
# ============================================================

# IMPORTANT:
# Do NOT add prefix="/admin" here.
#
# main.py already adds:
#     prefix="/admin"
#
# This prevents routes such as:
#     /admin/admin/timetable
#
# and gives:
#     /admin/timetable
#
router = APIRouter(
    tags=["Academic Management"],
)

user_timetable_router = APIRouter(
    prefix="/timetable",
    tags=["Timetable"],
)


# ============================================================
# ADMIN CHECK
# ============================================================

def admin_only(u: User):
    if u.role != "admin":
        raise HTTPException(
            status_code=403,
            detail="Admin access required.",
        )

    return u


# ============================================================
# PYDANTIC MODELS
# ============================================================

class CourseCreate(BaseModel):
    code: str
    title: str
    lecturer_id: int | None = None
    programme: str | None = None
    department: str | None = None


class CourseUpdate(BaseModel):
    title: str | None = None
    lecturer_id: int | None = None
    programme: str | None = None
    department: str | None = None
    active: bool | None = None


class EnrollmentCreate(BaseModel):
    student_id: int
    course_id: int


class TimetableCreate(BaseModel):
    course_id: int
    day: str
    start: str
    end: str
    room: str | None = None
    group: str | None = None
    block: str | None = None
    class_mode: str = "PHYSICAL"


# ============================================================
# TEXT HELPERS
# ============================================================

def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def normalize_text(value):
    value = clean_text(value)

    value = value.replace("\n", " ")
    value = value.replace("\r", " ")
    value = re.sub(r"\s+", " ", value)

    return value.strip().lower()


def normalize_course_code(value):
    value = clean_text(value)

    value = value.upper()

    value = value.replace(" ", "")
    value = value.replace("\n", "")
    value = value.replace("\r", "")

    return value


# ============================================================
# JSON SERIALIZERS
# ============================================================

def course_json(course: Course):
    lecturer_name = None

    if course.lecturer:
        lecturer_name = course.lecturer.full_name

    return {
        "id": course.id,
        "code": course.code,
        "title": course.title,
        "programme": course.programme,
        "lecturer_id": course.lecturer_id,
        "lecturer_name": lecturer_name,
    }


def timetable_json(item: Timetable):
    lecturer_id = None
    lecturer_name = None

    if item.course:
        lecturer_id = item.course.lecturer_id

        if item.course.lecturer:
            lecturer_name = item.course.lecturer.full_name

    return {
        "id": item.id,

        "course_id": item.course_id,

        "course_code": (
            item.course.code
            if item.course
            else ""
        ),

        "course_title": (
            item.course.title
            if item.course
            else ""
        ),

        "programme": (
            item.course.programme
            if item.course
            else None
        ),

        "lecturer_id": lecturer_id,

        "lecturer_name": lecturer_name,

        "day": item.day,

        "start": item.start,

        "end": item.end,

        "room": item.room,

        "group": item.group_name,

        "block": getattr(
            item,
            "block",
            None,
        ),

        "class_mode": item.class_mode,

        "is_active": item.is_active,
    }


# ============================================================
# ADMIN DASHBOARD
# ============================================================

@router.get("/dashboard")
def admin_dashboard(
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    students = (
        db.query(User)
        .filter(User.role == "student")
        .count()
    )

    lecturers = (
        db.query(User)
        .filter(User.role == "lecturer")
        .count()
    )

    courses = db.query(Course).count()

    timetable_entries = (
        db.query(Timetable)
        .filter(Timetable.is_active == True)
        .count()
    )

    enrollments = db.query(Enrollment).count()

    return {
        "students": students,
        "lecturers": lecturers,
        "courses": courses,
        "timetable_entries": timetable_entries,
        "enrollments": enrollments,
    }


# ============================================================
# ADMIN - LECTURERS
# ============================================================

@router.get("/lecturers")
def get_lecturers(
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    lecturers = (
        db.query(User)
        .filter(User.role == "lecturer")
        .order_by(User.full_name)
        .all()
    )

    return [
        {
            "id": lecturer.id,
            "full_name": lecturer.full_name,
            "email": lecturer.email,
            "staff_id": lecturer.staff_id,
            "department": lecturer.department,
            "is_active": lecturer.is_active,
        }
        for lecturer in lecturers
    ]


# ============================================================
# ADMIN - STUDENTS
# ============================================================

@router.get("/students")
def get_students(
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    students = (
        db.query(User)
        .filter(User.role == "student")
        .order_by(User.full_name)
        .all()
    )

    return [
        {
            "id": student.id,
            "full_name": student.full_name,
            "email": student.email,
            "student_id": student.student_id,
            "programme": student.programme,
            "department": student.department,
            "is_active": student.is_active,
        }
        for student in students
    ]


# ============================================================
# ADMIN - CREATE COURSE
# ============================================================

@router.post("/courses")
def create_course(
    data: CourseCreate,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    code = normalize_course_code(data.code)

    if not code:
        raise HTTPException(
            status_code=400,
            detail="Course code is required.",
        )

    existing = (
        db.query(Course)
        .filter(Course.code == code)
        .first()
    )

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Course already exists.",
        )

    lecturer = None

    if data.lecturer_id is not None:
        lecturer = (
            db.query(User)
            .filter(
                User.id == data.lecturer_id,
                User.role == "lecturer",
            )
            .first()
        )

        if not lecturer:
            raise HTTPException(
                status_code=404,
                detail="Lecturer not found.",
            )

    course = Course(
        code=code,
        title=clean_text(data.title),
        lecturer_id=data.lecturer_id,
        programme=clean_text(data.programme)
        or None,
    )

    db.add(course)
    db.commit()
    db.refresh(course)

    return course_json(course)


# ============================================================
# ADMIN - GET COURSES
# ============================================================

@router.get("/courses")
def get_courses(
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    courses = (
        db.query(Course)
        .order_by(Course.code)
        .all()
    )

    return [
        course_json(course)
        for course in courses
    ]


# ============================================================
# ADMIN - UPDATE COURSE
# ============================================================

@router.patch("/courses/{course_id}")
def update_course(
    course_id: int,
    data: CourseUpdate,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    course = (
        db.query(Course)
        .filter(Course.id == course_id)
        .first()
    )

    if not course:
        raise HTTPException(
            status_code=404,
            detail="Course not found.",
        )

    if data.title is not None:
        course.title = clean_text(data.title)

    if data.programme is not None:
        course.programme = clean_text(
            data.programme
        )

    if data.lecturer_id is not None:
        lecturer = (
            db.query(User)
            .filter(
                User.id == data.lecturer_id,
                User.role == "lecturer",
            )
            .first()
        )

        if not lecturer:
            raise HTTPException(
                status_code=404,
                detail="Lecturer not found.",
            )

        course.lecturer_id = lecturer.id

    db.commit()
    db.refresh(course)

    return course_json(course)


# ============================================================
# ADMIN - ENROLL STUDENT
# ============================================================

@router.post("/enrollments")
def create_enrollment(
    data: EnrollmentCreate,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    student = (
        db.query(User)
        .filter(
            User.id == data.student_id,
            User.role == "student",
        )
        .first()
    )

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student not found.",
        )

    course = (
        db.query(Course)
        .filter(Course.id == data.course_id)
        .first()
    )

    if not course:
        raise HTTPException(
            status_code=404,
            detail="Course not found.",
        )

    existing = (
        db.query(Enrollment)
        .filter(
            Enrollment.student_id
            == data.student_id,
            Enrollment.course_id
            == data.course_id,
        )
        .first()
    )

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Student is already enrolled.",
        )

    enrollment = Enrollment(
        student_id=data.student_id,
        course_id=data.course_id,
    )

    db.add(enrollment)
    db.commit()
    db.refresh(enrollment)

    return {
        "message": "Student enrolled successfully.",
        "id": enrollment.id,
        "student_id": enrollment.student_id,
        "course_id": enrollment.course_id,
    }


# ============================================================
# ADMIN - REMOVE ENROLLMENT
# ============================================================

@router.delete(
    "/enrollments/{student_id}/{course_id}"
)
def delete_enrollment(
    student_id: int,
    course_id: int,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    enrollment = (
        db.query(Enrollment)
        .filter(
            Enrollment.student_id == student_id,
            Enrollment.course_id == course_id,
        )
        .first()
    )

    if not enrollment:
        raise HTTPException(
            status_code=404,
            detail="Enrollment not found.",
        )

    db.delete(enrollment)
    db.commit()

    return {
        "message": "Enrollment removed successfully."
    }


# ============================================================
# ADMIN - COURSE STUDENTS
# ============================================================

@router.get("/courses/{course_id}/students")
def course_students(
    course_id: int,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    course = (
        db.query(Course)
        .filter(Course.id == course_id)
        .first()
    )

    if not course:
        raise HTTPException(
            status_code=404,
            detail="Course not found.",
        )

    students = (
        db.query(User)
        .join(
            Enrollment,
            Enrollment.student_id == User.id,
        )
        .filter(
            Enrollment.course_id == course_id,
            User.role == "student",
        )
        .order_by(User.full_name)
        .all()
    )

    return [
        {
            "id": student.id,
            "full_name": student.full_name,
            "email": student.email,
            "student_id": student.student_id,
            "programme": student.programme,
        }
        for student in students
    ]


# ============================================================
# ADMIN - CREATE TIMETABLE ENTRY
# ============================================================

@router.post("/timetable")
def create_timetable(
    data: TimetableCreate,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    course = (
        db.query(Course)
        .filter(Course.id == data.course_id)
        .first()
    )

    if not course:
        raise HTTPException(
            status_code=404,
            detail="Course not found.",
        )

    day = clean_text(data.day).title()
    start = clean_text(data.start)
    end = clean_text(data.end)

    existing = (
        db.query(Timetable)
        .filter(
            Timetable.course_id == data.course_id,
            Timetable.day == day,
            Timetable.start == start,
            Timetable.end == end,
            Timetable.is_active == True,
        )
        .first()
    )

    if existing:
        raise HTTPException(
            status_code=409,
            detail="This timetable entry already exists.",
        )

    class_mode = (
        clean_text(data.class_mode)
        .upper()
        or "PHYSICAL"
    )

    timetable = Timetable(
        course_id=data.course_id,
        day=day,
        start=start,
        end=end,
        room=clean_text(data.room) or None,
        group_name=clean_text(data.group) or None,
        block=clean_text(data.block) or None,
        class_mode=class_mode,
        is_active=True,
    )

    db.add(timetable)
    db.commit()
    db.refresh(timetable)

    return timetable_json(timetable)


# ============================================================
# ADMIN - GET ALL TIMETABLE
# ============================================================

@router.get("/timetable")
def admin_timetable(
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    rows = (
        db.query(Timetable)
        .filter(Timetable.is_active == True)
        .order_by(
            Timetable.day,
            Timetable.start,
            Timetable.id,
        )
        .all()
    )

    return [
        timetable_json(item)
        for item in rows
    ]


# ============================================================
# ADMIN - DELETE TIMETABLE
# ============================================================

@router.delete("/timetable/{timetable_id}")
def delete_timetable(
    timetable_id: int,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    timetable = (
        db.query(Timetable)
        .filter(Timetable.id == timetable_id)
        .first()
    )

    if not timetable:
        raise HTTPException(
            status_code=404,
            detail="Timetable entry not found.",
        )

    timetable.is_active = False

    db.commit()

    return {
        "message": "Timetable entry removed successfully."
    }


# ============================================================
# EXCEL HELPERS
# ============================================================

def excel_time(value):
    """
    Converts timetable spreadsheet values into HH:MM.

    Examples:
        08-00  -> 08:00
        8-00   -> 08:00
        08:00  -> 08:00
        8:00   -> 08:00
    """

    if value is None:
        return None

    text = clean_text(value)

    match = re.search(
        r"(\d{1,2})\s*[:\-]\s*(\d{2})",
        text,
    )

    if not match:
        return None

    hour = int(match.group(1))
    minute = int(match.group(2))

    if hour > 23 or minute > 59:
        return None

    return f"{hour:02d}:{minute:02d}"


def split_course_codes(value):
    """
    Splits cells containing:

        ICT401
        ICT401/ICT402
        ICT401 | ICT402
        ICT401\\ICT402
    """

    text = clean_text(value)

    if not text:
        return []

    parts = re.split(
        r"[/|\\]+",
        text,
    )

    result = []

    for part in parts:
        part = clean_text(part)

        if part:
            result.append(part)

    return result


def find_lecturer(
    db: Session,
    lecturer_text: str,
):
    """
    Attempts to match the lecturer name
    from the timetable to a registered lecturer.
    """

    target = normalize_text(
        lecturer_text
    )

    if not target:
        return None

    lecturers = (
        db.query(User)
        .filter(User.role == "lecturer")
        .all()
    )

    # Exact match
    for lecturer in lecturers:
        if normalize_text(
            lecturer.full_name
        ) == target:
            return lecturer

    # Partial match
    for lecturer in lecturers:
        name = normalize_text(
            lecturer.full_name
        )

        if (
            target in name
            or name in target
        ):
            return lecturer

    # Compare words
    target_words = set(
        target.split()
    )

    if target_words:
        for lecturer in lecturers:
            lecturer_words = set(
                normalize_text(
                    lecturer.full_name
                ).split()
            )

            if (
                target_words
                and lecturer_words
                and len(
                    target_words
                    & lecturer_words
                )
                >= max(
                    1,
                    min(
                        len(target_words),
                        len(lecturer_words),
                    ) // 2,
                )
            ):
                return lecturer

    return None


def find_or_create_course(
    db: Session,
    course_code: str,
    lecturer=None,
):
    """
    Finds an existing course or creates one.

    For this thesis prototype, uploaded timetable
    courses are treated as ICT courses.
    """

    code = normalize_course_code(
        course_code
    )

    if not code:
        return None

    course = (
        db.query(Course)
        .filter(Course.code == code)
        .first()
    )

    if course:
        changed = False

        # Prototype timetable is ICT
        if not course.programme:
            course.programme = "ICT"
            changed = True

        # If lecturer was successfully matched,
        # attach the lecturer to the course.
        if (
            lecturer
            and course.lecturer_id != lecturer.id
        ):
            course.lecturer_id = lecturer.id
            changed = True

        if changed:
            db.flush()

        return course

    course = Course(
        code=code,
        title=code,
        programme="ICT",
        lecturer_id=(
            lecturer.id
            if lecturer
            else None
        ),
    )

    db.add(course)
    db.flush()

    return course


# ============================================================
# ADMIN - UPLOAD EXCEL TIMETABLE
# ============================================================

@router.post("/timetable/upload")
async def upload_timetable(
    file: UploadFile = File(...),
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    admin_only(u)

    filename = (
        file.filename or ""
    ).lower()

    if not (
        filename.endswith(".xlsx")
        or filename.endswith(".xlsm")
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Only .xlsx or .xlsm "
                "Excel files are supported."
            ),
        )

    contents = await file.read()

    if not contents:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    try:
        workbook = load_workbook(
            filename=BytesIO(contents),
            data_only=True,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read the Excel file: "
                f"{exc}"
            ),
        )

    # ========================================================
    # SELECT SHEET
    # ========================================================

    preferred_sheet = (
        "Block_Teaching Main TimeTable"
    )

    if preferred_sheet in workbook.sheetnames:
        sheet = workbook[
            preferred_sheet
        ]
    else:
        sheet = workbook[
            workbook.sheetnames[0]
        ]

    created = 0
    skipped = 0
    courses_created = 0
    courses_updated = 0

    imported_items = []

    # ========================================================
    # PROCESS TIMETABLE
    # ========================================================

    row = 1

    while row <= sheet.max_row:

        day_value = sheet.cell(
            row=row,
            column=1,
        ).value

        day_text = clean_text(
            day_value
        )

        # ----------------------------------------------------
        # Ignore empty rows
        # ----------------------------------------------------

        if not day_text:
            row += 1
            continue

        normalized_day = normalize_text(
            day_text
        )

        day_map = {
            "mon": "Monday",
            "monday": "Monday",

            "tue": "Tuesday",
            "tues": "Tuesday",
            "tuesday": "Tuesday",

            "wed": "Wednesday",
            "wednesday": "Wednesday",

            "thu": "Thursday",
            "thur": "Thursday",
            "thurs": "Thursday",
            "thursday": "Thursday",

            "fri": "Friday",
            "friday": "Friday",

            "sat": "Saturday",
            "saturday": "Saturday",

            "sun": "Sunday",
            "sunday": "Sunday",
        }

        day = day_map.get(
            normalized_day
        )

        if not day:
            row += 1
            continue

        # ====================================================
        # TIME SLOTS
        # ====================================================

        # Columns 2-9 contain the timetable slots
        for col in range(
            2,
            min(
                sheet.max_column,
                9,
            ) + 1,
        ):

            start_value = sheet.cell(
                row=2,
                column=col,
            ).value

            end_value = sheet.cell(
                row=3,
                column=col,
            ).value

            start = excel_time(
                start_value
            )

            end = excel_time(
                end_value
            )

            if not start or not end:
                continue

            # ------------------------------------------------
            # Course row
            # ------------------------------------------------

            course_value = sheet.cell(
                row=row,
                column=col,
            ).value

            course_codes = (
                split_course_codes(
                    course_value
                )
            )

            if not course_codes:
                continue

            # ------------------------------------------------
            # Lecturer row
            # ------------------------------------------------

            lecturer_value = ""

            if row + 1 <= sheet.max_row:
                lecturer_value = clean_text(
                    sheet.cell(
                        row=row + 1,
                        column=col,
                    ).value
                )

            lecturer = find_lecturer(
                db,
                lecturer_value,
            )

            # ------------------------------------------------
            # Room row
            # ------------------------------------------------

            room_value = ""

            if row + 2 <= sheet.max_row:
                room_value = clean_text(
                    sheet.cell(
                        row=row + 2,
                        column=col,
                    ).value
                )

            # ------------------------------------------------
            # Block
            # ------------------------------------------------

            block_value = ""

            if sheet.max_column >= 10:
                block_value = clean_text(
                    sheet.cell(
                        row=row,
                        column=10,
                    ).value
                )

            # ------------------------------------------------
            # Class mode
            # ------------------------------------------------

            room_normalized = normalize_text(
                room_value
            )

            if any(
                word in room_normalized
                for word in [
                    "online",
                    "zoom",
                    "virtual",
                ]
            ):
                class_mode = "ONLINE"
            else:
                class_mode = "PHYSICAL"

            # =================================================
            # CREATE TIMETABLE FOR EACH COURSE
            # =================================================

            for course_code in course_codes:

                before_course = (
                    db.query(Course)
                    .filter(
                        Course.code
                        == normalize_course_code(
                            course_code
                        )
                    )
                    .first()
                )

                course = find_or_create_course(
                    db=db,
                    course_code=course_code,
                    lecturer=lecturer,
                )

                if not course:
                    skipped += 1
                    continue

                if before_course is None:
                    courses_created += 1
                else:
                    courses_updated += 1

                # ------------------------------------------------
                # Check duplicate timetable
                # ------------------------------------------------

                existing = (
                    db.query(Timetable)
                    .filter(
                        Timetable.course_id
                        == course.id,

                        Timetable.day
                        == day,

                        Timetable.start
                        == start,

                        Timetable.end
                        == end,

                        Timetable.room
                        == (
                            room_value
                            or None
                        ),

                        Timetable.is_active
                        == True,
                    )
                    .first()
                )

                if existing:
                    skipped += 1
                    continue

                timetable = Timetable(
                    course_id=course.id,
                    day=day,
                    start=start,
                    end=end,
                    room=(
                        room_value
                        or None
                    ),
                    group_name=None,
                    block=(
                        block_value
                        or None
                    ),
                    class_mode=class_mode,
                    is_active=True,
                )

                db.add(timetable)

                imported_items.append(
                    {
                        "course_code":
                            course.code,

                        "course_title":
                            course.title,

                        "lecturer":
                            (
                                lecturer.full_name
                                if lecturer
                                else None
                            ),

                        "day": day,

                        "start": start,

                        "end": end,

                        "room":
                            room_value or None,

                        "block":
                            block_value or None,

                        "class_mode":
                            class_mode,
                    }
                )

                created += 1

        # Move to the next 3-row timetable block
        row += 3

    # ========================================================
    # SAVE DATABASE
    # ========================================================

    try:
        db.commit()

    except Exception as exc:
        db.rollback()

        raise HTTPException(
            status_code=500,
            detail=(
                "Failed to save timetable: "
                f"{exc}"
            ),
        )

    return {
        "message": "Timetable uploaded successfully.",

        "filename": file.filename,

        "sheet": sheet.title,

        "created": created,

        "skipped": skipped,

        "courses_created":
            courses_created,

        "courses_updated":
            courses_updated,

        "total_imported":
            len(imported_items),

        "prototype_mode": True,

        "programme": "ICT",

        "data": imported_items,
    }


# ============================================================
# USER TIMETABLE
# ============================================================

@user_timetable_router.get("")
def user_timetable(
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    """
    Returns timetable according to the logged-in user.

    ADMIN:
        All active timetable entries.

    LECTURER:
        Only timetable entries for courses assigned
        to that lecturer.

        IMPORTANT:
        This returns the COMPLETE timetable.
        It does NOT filter by today's day.

    STUDENT:
        Prototype mode:
        All active timetable entries from the uploaded
        ICT timetable.

        Enrollment is intentionally NOT required.
    """

    # ========================================================
    # ADMIN
    # ========================================================

    if u.role == "admin":

        rows = (
            db.query(Timetable)
            .filter(
                Timetable.is_active == True
            )
            .order_by(
                Timetable.day,
                Timetable.start,
                Timetable.id,
            )
            .all()
        )

        return [
            timetable_json(item)
            for item in rows
        ]

    # ========================================================
    # LECTURER
    # ========================================================

    if u.role == "lecturer":

        rows = (
            db.query(Timetable)
            .join(
                Course,
                Timetable.course_id
                == Course.id,
            )
            .filter(
                Course.lecturer_id == u.id,
                Timetable.is_active == True,
            )
            .order_by(
                Timetable.day,
                Timetable.start,
                Timetable.id,
            )
            .all()
        )

        return [
            timetable_json(item)
            for item in rows
        ]

    # ========================================================
    # STUDENT
    # ========================================================

    if u.role == "student":

        # ====================================================
        # IMPORTANT PROTOTYPE CHANGE
        # ====================================================
        #
        # Students do NOT need Enrollment records.
        #
        # The currently uploaded timetable is the ICT
        # timetable for this thesis prototype.
        #
        # Therefore every student can retrieve the active
        # timetable.
        #
        # Enrollment can still be used later when the system
        # becomes a production university system.
        # ====================================================

        rows = (
            db.query(Timetable)
            .join(
                Course,
                Timetable.course_id
                == Course.id,
            )
            .filter(
                Timetable.is_active == True
            )
            .order_by(
                Timetable.day,
                Timetable.start,
                Timetable.id,
            )
            .all()
        )

        return [
            timetable_json(item)
            for item in rows
        ]

    raise HTTPException(
        status_code=403,
        detail="Unsupported user role.",
    )


# ============================================================
# USER TIMETABLE BY DAY
# ============================================================

@user_timetable_router.get(
    "/day/{day_name}"
)
def user_timetable_by_day(
    day_name: str,
    u: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    """
    Returns timetable entries for a particular day.

    Example:
        GET /timetable/day/Monday
    """

    day = clean_text(
        day_name
    ).title()

    # ========================================================
    # ADMIN
    # ========================================================

    if u.role == "admin":

        rows = (
            db.query(Timetable)
            .filter(
                Timetable.day == day,
                Timetable.is_active == True,
            )
            .order_by(
                Timetable.start,
                Timetable.id,
            )
            .all()
        )

        return [
            timetable_json(item)
            for item in rows
        ]

    # ========================================================
    # LECTURER
    # ========================================================

    if u.role == "lecturer":

        rows = (
            db.query(Timetable)
            .join(
                Course,
                Timetable.course_id
                == Course.id,
            )
            .filter(
                Course.lecturer_id == u.id,
                Timetable.day == day,
                Timetable.is_active == True,
            )
            .order_by(
                Timetable.start,
                Timetable.id,
            )
            .all()
        )

        return [
            timetable_json(item)
            for item in rows
        ]

    # ========================================================
    # STUDENT
    # ========================================================

    if u.role == "student":

        # Prototype:
        # No enrollment required.
        #
        # The uploaded timetable is ICT.

        rows = (
            db.query(Timetable)
            .join(
                Course,
                Timetable.course_id
                == Course.id,
            )
            .filter(
                Timetable.day == day,
                Timetable.is_active == True,
            )
            .order_by(
                Timetable.start,
                Timetable.id,
            )
            .all()
        )

        return [
            timetable_json(item)
            for item in rows
        ]

    raise HTTPException(
        status_code=403,
        detail="Unsupported user role.",
    )